import csv, io, os, re, sys, zipfile

CHECK_BELOW = 5
TEXT_EXT = {".csv", ".tsv", ".txt", ".dat", ".psv"}
EXCEL_EXT = {".xlsx", ".xlsm", ".xls"}
_NUM = re.compile(r"^[-+(]?[$€£]?\d[\d,]*(\.\d+)?([eE][-+]?\d+)?%?\)?$|^[-+]?\.\d+$")
_DATE = re.compile(r"^\d{4}[-/.]\d{1,2}([-/.]\d{1,2})?([ T]\d{1,2}:\d{2}(:\d{2})?)?$|^\d{1,2}[-/.]\d{1,2}[-/.]\d{2,4}$|"
                   r"^\d{1,2}-[A-Za-z]{3}-\d{2,4}$|^[A-Za-z]{3,9} \d{1,2},? \d{4}$")
_KV = re.compile(r"^\s*[A-Za-z][\w .()/&%-]{0,40}:(\s|$)")
_COMMENT = re.compile(r"^\s*(#|%|//|;)")
_UNIT = re.compile(r"^\s*[(\[]?\s*(in|units?|amounts? in|values? in|all figures in|currency)\b", re.I)


def cell_type(v) -> str:
    if v is None:
        return "empty"
    if isinstance(v, bool):
        return "text"
    if isinstance(v, (int, float)):
        return "num"
    if hasattr(v, "year"):
        return "date"
    s = str(v).strip()
    if not s:
        return "empty"
    if _NUM.match(s) or s.lower() in ("nan", "na", "n/a", "null", "-", "--"):
        return "num"
    if _DATE.match(s):
        return "date"
    return "text"


def _cells(row):
    return [c for c in row if cell_type(c) != "empty"]


def line_kind(row, raw="") -> str:
    cells = _cells(row)
    if not cells:
        return "blank"
    if raw and _COMMENT.match(raw):
        return "comment"
    first = str(cells[0])
    if isinstance(cells[0], str) and (_KV.match(first) or first.rstrip().endswith(":")) and len(cells) <= 3:
        return "keyvalue"
    if len(cells) == 1 and isinstance(cells[0], str) and _UNIT.match(first):
        return "unit"
    if len(cells) == 1 and isinstance(cells[0], str) and (len(first) > 40 or len(first.split()) >= 6):
        return "prose"
    return "row"


class Rows:

    def __init__(self, path, data=None, sheet=None):
        self.path, self.rows, self.raws, self.merged, self.done = path, [], [], [], False
        ext = os.path.splitext(path)[1].lower()
        if ext in EXCEL_EXT:
            import openpyxl
            src = io.BytesIO(data) if data is not None else path
            wb = openpyxl.load_workbook(src, read_only=False, data_only=True)
            ws = wb[sheet] if sheet else wb.worksheets[0]
            self.merged = [(r.min_row - 1, r.max_row - 1, r.min_col - 1, r.max_col - 1) for r in ws.merged_cells.ranges]
            self._it = (list(r) for r in ws.iter_rows(values_only=True))
            self.delim = None
        else:
            self._f = open(path, newline="", errors="replace")
            head = self._f.read(64 * 1024)
            self._f.seek(0)
            self.delim = _sniff(head)
            self._it = self._text_iter()

    def _text_iter(self):
        for line in self._f:
            raw = line.rstrip("\r\n")
            if self.delim == "ws":
                row = re.split(r"\s{2,}|\t", raw.strip()) if raw.strip() else []
            else:
                row = next(csv.reader([raw], delimiter=self.delim), [])
            yield row, raw

    def get(self, i):
        while len(self.rows) <= i and not self.done:
            try:
                x = next(self._it)
            except StopIteration:
                self.done = True
                break
            row, raw = x if self.delim is not None else (x, "")
            self.rows.append(row)
            self.raws.append(raw)
        return self.rows[i] if i < len(self.rows) else None

    def kind(self, i):
        r = self.get(i)
        return None if r is None else line_kind(r, self.raws[i])


def _sniff(head):
    lines = head.splitlines()[:200]
    sample = "\n".join(lines)
    try:
        return csv.Sniffer().sniff(sample, delimiters=",\t;|").delimiter
    except csv.Error:
        pass
    for d in ("\t", ",", ";", "|"):
        if sum(d in l for l in lines) >= max(2, len(lines) // 2):
            return d
    return "ws"


def _width(row):
    n = len(row)
    while n and cell_type(row[n - 1]) == "empty":
        n -= 1
    return n


def _col_types(rows, width):
    return [[cell_type(r[k]) if k < len(r) else "empty" for r in rows] for k in range(width)]


def _consistent(rows, width):
    cols = _col_types(rows, width)
    return sum(len({t for t in ts if t != "empty"}) <= 1 for ts in cols) >= 0.6 * max(1, len(cols)), cols


def _covered(i, upper, merged):
    ok = set()
    for k, v in enumerate(upper):
        if cell_type(v) == "empty":
            ok.add(k)
    for r0, r1, c0, c1 in merged:
        if r0 <= i - 1 <= r1:
            ok.update(range(c0, c1 + 1))
    for k in range(len(upper)):
        if cell_type(upper[k]) != "empty" and k + 1 < len(upper) and cell_type(upper[k + 1]) == "empty":
            ok.add(k)
    return ok


def detect(R: Rows) -> dict:
    i = 0
    while True:
        row = R.get(i)
        if row is None:
            return _no_header(R)
        if R.kind(i) != "row":
            i += 1
            continue
        cells = _cells(row)
        textish = sum(cell_type(c) == "text" for c in cells)
        if textish < 0.8 * len(cells) or not all(len(str(c)) <= 80 for c in cells):
            i += 1
            continue
        single = len(cells) == 1 and _width(row) == 1
        hw = _width(row)
        cont, j = [], i + 1
        while True:
            nxt = R.get(j)
            if nxt is None or R.kind(j) != "row":
                break
            nc = _cells(nxt)
            placed = {k for k, v in enumerate(nxt) if cell_type(v) != "empty"}
            upper = R.get(j - 1)
            if (len(nc) >= 1 and all(cell_type(c) == "text" for c in nc) and all(len(str(c)) <= 40 for c in nc)
                    and placed <= _covered(j, upper, R.merged)
                    and any(cell_type(v) != "empty" and k not in placed for k, v in enumerate(upper))):
                cont.append(j)
                j += 1
            else:
                break
        data, k = [], j
        while len(data) < CHECK_BELOW:
            r = R.get(k)
            if r is None:
                break
            kd = R.kind(k)
            if kd == "row":
                data.append(k)
            elif kd != "blank":
                break
            k += 1
        if len(data) < min(CHECK_BELOW, 2):
            i += 1
            continue
        drows = [R.get(x) for x in data]
        widths = [_width(r) for r in drows]
        named = {k for k, v in enumerate(row) if cell_type(v) != "empty"}
        for c in cont:
            named |= {k for k, v in enumerate(R.get(c)) if cell_type(v) != "empty"}
        if cont:
            hw = max(hw, max(named) + 1)
        if max(widths) > hw + max(1, hw // 10) or hw < 0.8 * max(widths) or len(named) < 0.8 * max(widths):
            i += 1
            continue
        ok, cols = _consistent(drows, hw)
        typed = any(t in ("num", "date") for ts in cols for t in ts)
        if not ok or (not typed and (hw < 3 or len({tuple(r) for r in drows}) < 2)) or (single and (not typed or max(widths) != 1)):
            i += 1
            continue
        if len(cells) == 1 and not single:
            i += 1
            continue
        ref = [next((t for t in ts if t != "empty"), "empty") for ts in cols]
        end, footer = _data_end(R, data[0], hw, ref)
        return {"header": i, "continuation": cont, "data_start": data[0], "data_end": end, "footer": footer,
                "width": hw, "typed": typed, "rows_read": len(R.rows),
                "skipped": [f"r{x}={R.kind(x)}" for x in range(i) if R.kind(x) != "blank"],
                "reason": (f"first row of short text filling the width of the {len(data)} rows below it, which keep "
                           f"{hw} fields and one type per column" + (f"; rows {cont} sit under merged or empty cells of the "
                           f"row above and hold only text, so they continue the header" if cont else ""))}


def _no_header(R):
    body = [k for k in range(len(R.rows)) if R.kind(k) == "row"][:CHECK_BELOW]
    if len(body) >= 2 and len({_width(R.get(k)) for k in body}) == 1 and all(
            any(cell_type(c) in ("num", "date") for c in R.get(k)) for k in body):
        return {"header": None, "no_header": True, "data_start": body[0], "rows_read": len(R.rows),
                "reason": "the first rows are already typed data of one shape; no header row"}
    return {"header": None, "reason": "no table region found", "rows_read": len(R.rows)}


def _conforms(row, kind, hw, ref):
    if kind != "row" or _width(row) > hw + max(1, hw // 10):
        return False
    same = tot = 0
    for k, t in enumerate(ref[:hw]):
        v = row[k] if k < len(row) else None
        ct = cell_type(v)
        if ct == "empty" or t == "empty":
            continue
        tot += 1
        same += ct == t
    return tot > 0 and same >= 0.6 * tot


def _data_end(R, start, hw, ref):
    last, k = start, start
    while True:
        r = R.get(k)
        if r is None:
            break
        if _conforms(r, R.kind(k), hw, ref):
            last = k
        k += 1
    footer = [x for x in range(last + 1, len(R.rows)) if R.kind(x) != "blank"]
    return last, footer


def combined_names(R: Rows, res: dict) -> list:
    h, cont, w = res["header"], res["continuation"], res["width"]
    top = [R.get(h)[k] if k < len(R.get(h)) else None for k in range(w)]
    filled = list(top)
    for r0, r1, c0, c1 in R.merged:
        if r0 <= h <= r1:
            for k in range(c0, min(c1 + 1, w)):
                filled[k] = top[c0]
    for c in cont:
        low = R.get(c)
        for k in range(1, w):
            if cell_type(filled[k]) == "empty" and k < len(low) and cell_type(low[k]) != "empty" and cell_type(filled[k - 1]) != "empty":
                filled[k] = filled[k - 1]
    names = [("" if cell_type(v) == "empty" else str(v).strip()) for v in filled]
    for c in cont:
        low = R.get(c)
        names = [(f"{n} | {str(low[k]).strip()}" if n else str(low[k]).strip()) if k < len(low) and cell_type(low[k]) != "empty" else n
                 for k, n in enumerate(names)]
    return names


def pandas_calls(path, res, delim, sheet=None):
    h, cont, de = res["header"], res["continuation"], res["data_end"]
    n = de - h - len(cont)
    if delim is None:
        src = (f"zipfile.ZipFile({path.split('!', 1)[0]!r}).open({path.split('!', 1)[1]!r})" if "!" in path else repr(path))
        base = f"pd.read_excel({src}" + (f", sheet_name={sheet!r}" if sheet else "")
        one = f"{base}, header={h}" + (f", skiprows={cont}" if cont else "") + f", nrows={n})"
        two = f"{base}, header=[{h}, {h + len(cont)}], nrows={n})" if cont else None
        if "!" in path:
            one, two = "import zipfile; " + one, two and "import zipfile; " + two
        return one, two
    sep = {"\t": ", sep='\\t'", ",": "", ";": ", sep=';'", "|": ", sep='|'", "ws": r", sep=r'\s{2,}', engine='python'"}[delim]
    skip = f"skiprows={list(range(h)) + cont}" if cont else f"skiprows={h}"
    one = f"pd.read_csv({path!r}{sep}, {skip}, header=0, nrows={n}, skip_blank_lines=False)"
    two = (f"pd.read_csv({path!r}{sep}, skiprows={h}, header=[0, {len(cont)}], nrows={n}, skip_blank_lines=False)"
           if cont else None)
    return one, two


def analyse(path, data=None, sheet=None) -> dict:
    R = Rows(path, data, sheet)
    res = detect(R)
    res.update(file=path, sheet=sheet, delim=R.delim)
    if res["header"] is None:
        return res
    res["names"] = combined_names(R, res) if res["continuation"] else None
    res["call_one"], res["call_two"] = pandas_calls(path, res, R.delim, sheet)
    return res


def analyse_file(path) -> list:
    ext = os.path.splitext(path)[1].lower()
    if ext in EXCEL_EXT:
        import openpyxl
        wb = openpyxl.load_workbook(path, read_only=True)
        names = wb.sheetnames
        wb.close()
        return [analyse(path, sheet=s) for s in names]
    if ext == ".zip":
        out = []
        with zipfile.ZipFile(path) as z:
            for n in z.namelist():
                if os.path.splitext(n)[1].lower() in EXCEL_EXT:
                    import openpyxl
                    blob = z.read(n)
                    wb = openpyxl.load_workbook(io.BytesIO(blob), read_only=True)
                    for s in wb.sheetnames:
                        out.append(analyse(f"{path}!{n}", data=blob, sheet=s))
                    wb.close()
        return out
    if ext in TEXT_EXT:
        return [analyse(path)]
    return []


def _v7_header(path, sheet):
    try:
        sys.path.insert(0, os.environ.get("PEEK_HELPERS_DIR", os.path.dirname(os.path.abspath(__file__))))
        import openpyxl, peek_data as pk
        wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
        rows = {k: pk.trimmed(r) for k, r in enumerate(wb[sheet].iter_rows(values_only=True)) if pk.filled(r)}
        wb.close()
        return pk.find_header(rows)[0]
    except Exception:
        return None


def preview_lines(path: str) -> list[str]:
    lines = []
    for r in analyse_file(path):
        tag = f"sheet {r['sheet']}: " if r.get("sheet") else ""
        if r["header"] is None:
            if r.get("no_header"):
                sep = {"\t": ", sep='\\t'", ",": "", ";": ", sep=';'", "|": ", sep='|'", "ws": r", sep=r'\s{2,}', engine='python'"}.get(r["delim"], "")
                skip = f", skiprows={r['data_start']}" if r["data_start"] else ""
                lines.append(f"{tag}no header row: the first rows are already data; read with pd.read_csv({path!r}{sep}{skip}, header=None)")
            continue
        plain = r["header"] == 0 and not r["continuation"] and not r["footer"]
        if not plain:
            above = ", ".join(r["skipped"][:8]) or "none"
            lines.append(f"{tag}header at physical row {r['header']}"
                         + (f" with continuation row(s) {r['continuation']}" if r["continuation"] else "")
                         + f"; data rows {r['data_start']}-{r['data_end']}"
                         + (f"; {len(r['footer'])} footer row(s) after the data" if r["footer"] else "")
                         + f" (rows above: {above}). Why: {r['reason']}.")
            lines.append(f"{tag}read as a single-row header" + (" (the continuation rows skipped)" if r["continuation"] else "")
                         + f": {r['call_one']}")
            if r["call_two"]:
                lines.append(f"{tag}read as a two-row header: {r['call_two']}; or read with the single-row call and set "
                             f"df.columns = {r['names']!r}")
        if r.get("sheet") and os.path.splitext(path)[1].lower() in EXCEL_EXT:
            v7 = _v7_header(path, r["sheet"])
            if v7 is not None and v7 != r["header"]:
                lines.append(f"{tag}the other header rule picks row {v7}; both candidates are listed, check which one the data needs")
    return lines
