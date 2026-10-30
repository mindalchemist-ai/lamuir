import csv, io, os, re, sys, zipfile

SCAN_ROWS = 60
CHECK_BELOW = 5
TEXT_EXT = {".csv", ".tsv", ".txt", ".dat", ".psv"}
EXCEL_EXT = {".xlsx", ".xlsm", ".xls"}
_NUM = re.compile(r"^[-+(]?\$?\d[\d,]*(\.\d+)?([eE][-+]?\d+)?%?\)?$|^[-+]?\.\d+$")
_DATE = re.compile(r"^\d{4}[-/.]\d{1,2}([-/.]\d{1,2})?([ T]\d{1,2}:\d{2}(:\d{2})?)?|^\d{1,2}[-/.]\d{1,2}[-/.]\d{2,4}$|"
                   r"^\d{1,2}-[A-Za-z]{3}-\d{2,4}$|^[A-Za-z]{3,9} \d{1,2},? \d{4}$")
_KV = re.compile(r"^\s*[A-Za-z][\w .()/&%-]{0,40}:\s*\S")
_COMMENT = re.compile(r"^\s*(#|%|//|;)")


def cell_type(v) -> str:
    if v is None:
        return "empty"
    if isinstance(v, bool):
        return "text"
    if isinstance(v, (int, float)):
        return "num"
    if hasattr(v, "year"):
        return "date"
    s = str(v).strip()
    if not s:
        return "empty"
    if _NUM.match(s) or s.lower() in ("nan", "na", "n/a", "null", "-"):
        return "num"
    if _DATE.match(s):
        return "date"
    return "text"


def _trim(row):
    row = list(row)
    while row and cell_type(row[-1]) == "empty":
        row.pop()
    return row


def line_kind(row, raw: str = "") -> str:
    cells = [c for c in row if cell_type(c) != "empty"]
    if not cells:
        return "blank"
    if raw and _COMMENT.match(raw):
        return "comment"
    if len(cells) <= 2 and isinstance(cells[0], str) and (_KV.match(cells[0]) or (len(cells) == 2 and str(cells[0]).rstrip().endswith(":"))):
        return "keyvalue"
    if len(cells) == 1 and isinstance(cells[0], str) and (len(cells[0]) > 40 or len(cells[0].split()) >= 6):
        return "prose"
    return "row"


def _sniff(text: str):
    sample = "\n".join(text.splitlines()[:SCAN_ROWS])
    try:
        return csv.Sniffer().sniff(sample, delimiters=",\t;|").delimiter
    except csv.Error:
        pass
    for d in ("\t", ",", ";", "|"):
        if sum(d in l for l in sample.splitlines()) >= max(2, len(sample.splitlines()) // 2):
            return d
    return "ws"


def text_rows(text: str):
    lines = text.splitlines()[:SCAN_ROWS]
    d = _sniff(text)
    if d == "ws":
        rows = [re.split(r"\s{2,}|\t", l.strip()) if l.strip() else [] for l in lines]
    else:
        rows = [next(csv.reader([l], delimiter=d), []) for l in lines]
    return [list(r) for r in rows], lines, d


def detect_rows(rows: list, raws: list = None) -> dict:
    raws = raws or [""] * len(rows)
    kinds = [line_kind(r, raws[i] if i < len(raws) else "") for i, r in enumerate(rows)]
    body = [i for i, k in enumerate(kinds) if k == "row"]
    widths = [len(rows[i]) for i in body]
    width = max(set(widths), key=widths.count) if widths else 0
    cands = []
    for n, i in enumerate(body):
        row = rows[i]
        cells = [c for c in row if cell_type(c) != "empty"]
        types = [cell_type(c) for c in cells]
        textish = sum(t == "text" for t in types)
        short = all(len(str(c)) <= 60 for c in cells)
        if not cells or textish < 0.8 * len(cells) or not short or len(cells) < 0.8 * width:
            cands.append((i, "not mostly short text filling the table width"))
            continue
        below = body[n + 1:n + 1 + CHECK_BELOW]
        if not below:
            cands.append((i, "no rows below"))
            continue
        cont = []
        for j in below:
            if sum(cell_type(c) == "text" for c in rows[j]) >= 2 and all(cell_type(c) in ("text", "empty") for c in rows[j]) and any(
                    cell_type(c) in ("num", "date") for k in body if k > j for c in rows[k][:width]):
                cont.append(j)
            else:
                break
        data = [j for j in body[n + 1 + len(cont):n + 1 + len(cont) + CHECK_BELOW]]
        if not data:
            cands.append((i, "no data rows below"))
            continue
        same_width = all(len(rows[j]) <= len(row) + max(1, len(row) // 10) for j in data)
        col_types = [[cell_type(rows[j][k]) if k < len(rows[j]) else "empty" for j in data] for k in range(len(row))]
        consistent = sum(len({t for t in ts if t != "empty"}) <= 1 for ts in col_types) >= 0.6 * len(col_types)
        typed = any(t in ("num", "date") for ts in col_types for t in ts)
        all_text = not typed
        if same_width and consistent and (typed or all_text):
            reason = (f"first row of short text ({textish}/{len(cells)} text cells) whose next {len(data)} rows have "
                      f"{len(row)} fields with one type per column")
            skipped = [f"r{k}={kinds[k]}" for k in range(i) if kinds[k] != "row"] + [f"r{k}=row" for k in body[:n]]
            return {"header": i, "continuation": cont, "width": len(row), "skipped": skipped, "reason": reason,
                    "candidates": cands + [(i, reason)], "doubt": "data rows are text only" if all_text else ""}
        cands.append((i, f"rows below: same width={same_width}, one type per column={consistent}"))
    first = body[:CHECK_BELOW]
    if first and all(any(cell_type(c) in ("num", "date") for c in rows[i]) for i in first) and len({len(rows[i]) for i in first}) == 1:
        return {"header": None, "no_header": True, "first_row": first[0], "continuation": [], "width": width,
                "skipped": [f"r{k}={kinds[k]}" for k in range(first[0])], "reason": "the first rows are already typed data; no header row",
                "candidates": cands, "doubt": ""}
    return {"header": None, "continuation": [], "width": width, "skipped": [], "reason": "no row met the rules",
            "candidates": cands, "doubt": ""}


def excel_header(path: str, data: bytes = None) -> list[dict]:
    sys.path.insert(0, os.environ.get("PEEK_HELPERS_DIR", os.path.dirname(os.path.abspath(__file__))))
    import peek_data as pk
    import openpyxl
    wb = openpyxl.load_workbook(io.BytesIO(data) if data else path, read_only=True, data_only=True)
    out = []
    for ws in wb.worksheets:
        rows = {}
        for k, r in enumerate(ws.iter_rows(values_only=True)):
            if k >= SCAN_ROWS:
                break
            if pk.filled(r):
                rows[k] = pk.trimmed(r)
        h, width = pk.find_header(rows)
        mine = detect_rows([list(rows.get(k, [])) for k in range(max(rows) + 1 if rows else 0)])
        out.append({"sheet": ws.title, "header": h, "width": width, "mine": mine["header"], "continuation": mine["continuation"]})
    wb.close()
    return out


def pandas_call(path: str, res: dict, delim: str) -> str:
    h = res["header"]
    sep = {"\t": ", sep='\\t'", ",": "", ";": ", sep=';'", "|": ", sep='|'", "ws": r", sep=r'\s{2,}', engine='python'"}[delim]
    if h is None:
        if res.get("no_header"):
            skip = f", skiprows={res['first_row']}" if res["first_row"] else ""
            return f"pd.read_csv({path!r}{sep}{skip}, header=None)"
        return ""
    hdr = f"header=list(range({len(res['continuation']) + 1}))" if res["continuation"] else "header=0"
    skip = f", skiprows={h}" if h else ""
    return f"pd.read_csv({path!r}{sep}{skip}, {hdr})"


def detect_file(path: str) -> list[dict]:
    ext = os.path.splitext(path)[1].lower()
    if ext in EXCEL_EXT:
        return [dict(file=path, kind="excel", **s) for s in excel_header(path)]
    if ext == ".zip":
        out = []
        with zipfile.ZipFile(path) as z:
            for n in z.namelist():
                e = os.path.splitext(n)[1].lower()
                if e in EXCEL_EXT:
                    out += [dict(file=f"{path}!{n}", kind="excel", **s) for s in excel_header(n, z.read(n))]
        return out
    if ext in TEXT_EXT:
        with open(path, errors="replace") as f:
            text = f.read(256 * 1024)
        rows, raws, d = text_rows(text)
        res = detect_rows(rows, raws)
        return [dict(file=path, kind="text", delim=d, call=pandas_call(path, res, d), **res)]
    return []


def header_name_notes(names) -> list[str]:
    out, seen = [], {}
    for k, n in enumerate(names):
        if n is None or str(n) == "":
            continue
        s = str(n)
        why = []
        if s != s.strip():
            why.append("leading or trailing spaces")
        if "\n" in s:
            why.append("a line break")
        sym = sorted(set(re.findall(r"[^\w\s()/%&.,:+-]", s)))
        if sym:
            why.append("the symbol(s) " + " ".join(sym))
        if why:
            out.append(f"column {k} is {s!r} ({', '.join(why)}); pandas keeps the name exactly so")
        seen.setdefault(s.strip(), []).append(k)
    for s, ks in seen.items():
        if len(ks) > 1:
            out.append(f"columns {', '.join(map(str, ks))} are all named {s!r}; pandas names the repeats {s!r}.1, {s!r}.2")
    return out


def preview_lines(path: str) -> list[str]:
    lines = []
    for r in detect_file(path):
        if r["kind"] == "text":
            if r.get("no_header"):
                lines.append(f"no header row: the first rows are already data; read with {r['call']}")
            elif r["header"]:
                above = ", ".join(r["skipped"]) or "none"
                lines.append(f"header is physical line {r['header']} (lines above: {above}); read with {r['call']}")
            if r["header"]:
                with open(path, errors="replace") as f:
                    rows, _, _ = text_rows(f.read(256 * 1024))
                names = rows[r["header"]] if r["header"] < len(rows) else []
                lines += [f"header note: {n}" for n in header_name_notes(names)]
        elif r["kind"] == "excel" and r["header"] is not None:
            import openpyxl
            wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
            ws = wb[r["sheet"]]
            names = next((list(row) for k, row in enumerate(ws.iter_rows(values_only=True)) if k == r["header"]), [])
            wb.close()
            lines += [f"sheet {r['sheet']}: header note: {n}" for n in header_name_notes(names)]
    return lines
